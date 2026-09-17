"""Office Open XML documents (docx, pptx, xlsx)."""

from __future__ import annotations

import zipfile

from app.config import settings
from app.extraction.base import (
    CorruptedFileError,
    ExtractionContext,
    ExtractionResult,
    LimitExceededError,
    PageResult,
)

MAX_XLSX_CELLS = 2_000_000


def _guard_zip(path: str) -> None:
    """OOXML files are zips: reject damaged archives and zip bombs before parsing."""
    try:
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
            total = sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile as exc:
        raise CorruptedFileError(f"The document archive is damaged: {exc}") from exc
    if bad:
        raise CorruptedFileError(f"The document archive is damaged (bad entry: {bad}).")
    if total > settings.max_archive_uncompressed_mb * 1024 * 1024:
        raise LimitExceededError(
            f"The document expands to {total / 1e6:.0f} MB; the limit is {settings.max_archive_uncompressed_mb} MB."
        )


def _core_properties(props) -> dict:
    fields = ("title", "author", "subject", "keywords", "last_modified_by", "created", "modified")
    out = {}
    for name in fields:
        value = getattr(props, name, None)
        out[name] = value.isoformat() if hasattr(value, "isoformat") else (value or None)
    return out


def extract_docx(path: str, ctx: ExtractionContext) -> ExtractionResult:
    _guard_zip(path)
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(path)
    except Exception as exc:
        raise CorruptedFileError(f"The Word document could not be parsed: {exc}") from exc

    blocks: list[str] = []
    tables = 0
    # Walk the body in order so tables stay where they appear in the document.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            blocks.append(Paragraph(child, document).text)
        elif tag == "tbl":
            tables += 1
            table = Table(child, document)
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = cell.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != value:  # merged cells repeat their value
                        cells.append(value)
                blocks.append("\t".join(cells))
            blocks.append("")
    for section in document.sections:
        header = "\n".join(p.text for p in section.header.paragraphs if p.text.strip())
        if header and header not in blocks:
            blocks.insert(0, header)

    text = "\n".join(blocks).strip()
    metadata = {
        "docx": {
            **_core_properties(document.core_properties),
            "paragraphs": len(document.paragraphs),
            "tables": tables,
            "images": len(document.inline_shapes),
        }
    }
    warnings = []
    if len(document.inline_shapes) and len(text) < 50:
        warnings.append("The document is mostly images; text inside images is not extracted from Word files.")
    return ExtractionResult(pages=[PageResult(1, text, "parser", len(text))], metadata=metadata, warnings=warnings)


def extract_pptx(path: str, ctx: ExtractionContext) -> ExtractionResult:
    _guard_zip(path)
    from pptx import Presentation

    try:
        prs = Presentation(path)
    except Exception as exc:
        raise CorruptedFileError(f"The presentation could not be parsed: {exc}") from exc

    def shape_text(shape) -> list[str]:
        out = []
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
            out.append(shape.text_frame.text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                out.append("\t".join(cell.text for cell in row.cells))
        for sub in getattr(shape, "shapes", []):  # grouped shapes
            out.extend(shape_text(sub))
        return out

    pages = []
    slides = list(prs.slides)
    for i, slide in enumerate(slides, start=1):
        ctx.check_cancelled()
        ctx.progress((i - 1) / max(len(slides), 1), "extracting", f"slide {i}/{len(slides)}")
        parts = []
        for shape in slide.shapes:
            parts.extend(shape_text(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[Notes] {notes}")
        text = "\n".join(parts).strip()
        pages.append(PageResult(i, text, "parser", len(text), warnings=[] if text else ["No text on slide"]))
    metadata = {"pptx": {**_core_properties(prs.core_properties), "slides": len(slides)}}
    return ExtractionResult(pages=pages or [PageResult(1, "", "parser")], metadata=metadata)


def extract_xlsx(path: str, ctx: ExtractionContext) -> ExtractionResult:
    _guard_zip(path)
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CorruptedFileError(f"The spreadsheet could not be parsed: {exc}") from exc

    pages, cells_seen, warnings = [], 0, []
    try:
        sheets = wb.worksheets
        for i, sheet in enumerate(sheets, start=1):
            ctx.check_cancelled()
            ctx.progress((i - 1) / max(len(sheets), 1), "extracting", f"sheet {i}/{len(sheets)}: {sheet.title}")
            lines = [f"# {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                cells_seen += len(row)
                if cells_seen > MAX_XLSX_CELLS:
                    warnings.append(f"Spreadsheet truncated after {MAX_XLSX_CELLS:,} cells.")
                    break
                values = ["" if v is None else str(v) for v in row]
                while values and values[-1] == "":
                    values.pop()
                if values:
                    lines.append("\t".join(values))
            text = "\n".join(lines)
            pages.append(PageResult(i, text, "parser", len(text)))
            if cells_seen > MAX_XLSX_CELLS:
                break
        metadata = {"xlsx": {**_core_properties(wb.properties), "sheets": [s.title for s in sheets]}}
    finally:
        wb.close()
    return ExtractionResult(pages=pages, metadata=metadata, warnings=warnings)
