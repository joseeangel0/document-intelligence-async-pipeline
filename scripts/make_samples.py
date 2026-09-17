"""Generate the sample documents used by the demo, the e2e tests and the chaos tests.

Run inside the app image (all libraries are there):
    docker compose run --rm -v "$PWD/samples:/srv/samples" --user root api python scripts/make_samples.py
"""

from __future__ import annotations

import io
import random
import zipfile
from pathlib import Path

import pymupdf
from PIL import Image, ImageFilter

OUT = Path(__file__).resolve().parent.parent / "samples"
OUT.mkdir(exist_ok=True)
EDGE = OUT / "edge_cases"
EDGE.mkdir(exist_ok=True)
random.seed(7)

INVOICE = """ACME Data Services S.A. de C.V.
Invoice No. INV-2026-0917
Date: 2026-09-17        Due date: 17/10/2026

Bill to: Universidad Tecnológica, Mérida, Yucatán
Contact: billing@universidad.example.mx   Phone: +52 999 123 4567

Description                                   Qty     Amount
Document OCR processing (pages)              1,200   $3,600.00 MXN
Storage (GB-month)                              50     $250.00 MXN
Priority support                                 1   $1,150.00 MXN

Subtotal: $5,000.00 MXN     IVA 16%: $800.00 MXN     Total: $5,800.00 MXN
Payment instructions: https://pay.example.com/inv/2026-0917
"""

PARAGRAPHS_EN = [
    "Real services rarely process a request synchronously, and file ingestion especially. Documents arrive "
    "faster than they can be processed, so a reliable solution accepts the upload, stores it, queues the work "
    "and lets users track progress asynchronously.",
    "The claim-check pattern keeps messages small: the heavy payload is stored in object storage and only a "
    "reference travels through the queue. Workers redeem the claim check when they are ready to do the work.",
    "Failure handling matters as much as the happy path. Workers can be killed in the middle of a task, brokers "
    "can restart and files can be corrupted. No job may be lost and no document may stay pending forever.",
]
PARAGRAPHS_ES = [
    "La digitalización de documentos permite que la información sea buscable y reutilizable. ¿Cuántas páginas "
    "escaneadas existen todavía en los archivos de una universidad? Probablemente millones.",
    "El reconocimiento óptico de caracteres convierte imágenes en texto. La precisión depende de la resolución, "
    "del contraste y de la tipografía; por eso conviene medir antes de elegir una herramienta. ¡Midamos!",
    "Año tras año, el volumen de facturas, contratos y constancias crece. Un sistema asíncrono evita que el "
    "usuario espere con la conexión abierta mientras el servidor procesa archivos pesados.",
]


def text_pdf(pages: list[str], fontsize: float = 11) -> pymupdf.Document:
    doc = pymupdf.open()
    for body in pages:
        page = doc.new_page(width=612, height=792)
        rect = pymupdf.Rect(60, 60, 552, 740)
        page.insert_textbox(rect, body, fontsize=fontsize, fontname="helv")
    return doc


def rasterize(doc: pymupdf.Document, dpi: int = 200) -> list[Image.Image]:
    images = []
    for page in doc:
        pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
        images.append(Image.frombytes("L", (pix.width, pix.height), pix.samples))
    return images


def scan_effect(img: Image.Image, angle: float = 0.8, noise: int = 18) -> Image.Image:
    img = img.rotate(angle, expand=True, fillcolor=255, resample=Image.Resampling.BICUBIC)
    px = img.load()
    for _ in range(img.width * img.height // 60):
        x, y = random.randrange(img.width), random.randrange(img.height)
        px[x, y] = max(0, min(255, px[x, y] + random.randint(-noise * 4, noise)))
    return img.filter(ImageFilter.GaussianBlur(0.6))


def image_pdf(images: list[Image.Image]) -> pymupdf.Document:
    doc = pymupdf.open()
    for img in images:
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=70)
        page = doc.new_page(width=612, height=792)
        page.insert_image(page.rect, stream=buf.getvalue())
    return doc


CONTRACT_SECTIONS = [
    ("1. Scope of services", "The provider will digitise and index the client's archive of invoices and contracts. "
     "All processing runs on premises; no document leaves the client's network."),
    ("2. Payment terms", "Invoices are payable within 30 days. Late payments accrue interest of 1.5% per month."),
]
PRICE_ROWS = [["Service", "Unit", "Price (MXN)"], ["OCR page", "page", "3.00"], ["Storage", "GB-month", "5.00"],
              ["Support", "hour", "1,150.00"]]


def structured_pdf() -> None:
    """Born-digital PDF with a title, headings, a bullet list and a ruled table (ground truth for Markdown)."""
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    y = 72
    page.insert_text((72, y), "Service Agreement", fontsize=22, fontname="hebo")
    y += 40
    for title, body in CONTRACT_SECTIONS:
        page.insert_text((72, y), title, fontsize=15, fontname="hebo")
        y += 10
        rect = pymupdf.Rect(72, y, 540, y + 50)
        page.insert_textbox(rect, body, fontsize=11, fontname="helv")
        y += 60
    page.insert_text((72, y), "3. Deliverables", fontsize=15, fontname="hebo")
    y += 22
    for item in ("Searchable text for every page", "Markdown with tables preserved", "Retrieval chunks with page numbers"):
        page.insert_text((84, y), f"\u2022 {item}", fontsize=11, fontname="helv")
        y += 18
    y += 20
    page.insert_text((72, y), "4. Price list", fontsize=15, fontname="hebo")
    y += 14
    col_x = [72, 250, 380, 520]
    for r, row in enumerate(PRICE_ROWS):
        top = y + r * 22
        for c, value in enumerate(row):
            page.draw_rect(pymupdf.Rect(col_x[c], top, col_x[c + 1], top + 22), color=(0, 0, 0), width=0.8)
            page.insert_text((col_x[c] + 6, top + 15), value, fontsize=11, fontname="hebo" if r == 0 else "helv")
    doc.set_metadata({"title": "Service Agreement", "author": "ACME Data Services"})
    doc.save(OUT / "contract_structured.pdf")


def scanned_structured() -> None:
    """Image-only scan of the structured contract: tables can only be recovered by layout analysis."""
    scan = scan_effect(rasterize(pymupdf.open(OUT / "contract_structured.pdf"), dpi=200)[0], angle=0.4, noise=10)
    image_pdf([scan]).save(OUT / "contract_scanned.pdf")


def structured_docx(scan: Image.Image) -> None:
    """Word file with heading styles, a bullet list, a table and a scanned invoice pasted as a picture."""
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading("Service Agreement", 0)
    for title, body in CONTRACT_SECTIONS:
        d.add_heading(title, 1)
        d.add_paragraph(body)
    d.add_heading("3. Deliverables", 1)
    for item in ("Searchable text for every page", "Markdown with tables preserved"):
        d.add_paragraph(item, style="List Bullet")
    d.add_heading("4. Price list", 1)
    table = d.add_table(rows=len(PRICE_ROWS), cols=3)
    for r, row in enumerate(PRICE_ROWS):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.add_heading("Annex A. Scanned invoice", 1)
    buf = io.BytesIO()
    scan.convert("L").save(buf, format="PNG")
    buf.seek(0)
    d.add_picture(buf, width=Inches(6))
    d.save(OUT / "contract_with_scan.docx")


def email_with_attachments() -> None:
    from email.message import EmailMessage

    msg = EmailMessage()
    msg["From"] = "Billing <billing@acme.example.com>"
    msg["To"] = "accounts@universidad.example.mx"
    msg["Subject"] = "Factura INV-2026-0917 y contrato"
    msg["Date"] = "Thu, 17 Sep 2026 10:30:00 -0600"
    msg.set_content("Hola,\n\nAdjuntamos la factura y el contrato firmado.\n\nSaludos")
    msg.add_alternative(
        "<p>Hola,</p><p>Adjuntamos la <b>factura</b> y el contrato firmado.</p>"
        "<table><tr><th>Documento</th><th>Total</th></tr><tr><td>INV-2026-0917</td><td>$5,800.00 MXN</td></tr></table>"
        "<p>Saludos</p>", subtype="html")
    msg.add_attachment((OUT / "invoice_digital.pdf").read_bytes(), maintype="application", subtype="pdf",
                       filename="invoice_digital.pdf")
    msg.add_attachment((OUT / "invoice_scan.png").read_bytes(), maintype="image", subtype="png",
                       filename="invoice_scan.png")
    msg.add_attachment(b"MZ\x90\x00" + bytes(100), maintype="application", subtype="octet-stream",
                       filename="tool.exe")
    (OUT / "email_invoice.eml").write_bytes(bytes(msg))


def main() -> None:
    # --- PDFs ---------------------------------------------------------------------------
    doc = text_pdf([INVOICE, "\n\n".join(PARAGRAPHS_EN + PARAGRAPHS_ES)])
    doc.set_metadata({"title": "Invoice INV-2026-0917", "author": "ACME Data Services", "subject": "Sample invoice"})
    doc.save(OUT / "invoice_digital.pdf")

    letter = text_pdf(["\n\n".join(PARAGRAPHS_ES), "\n\n".join(PARAGRAPHS_EN), INVOICE])
    image_pdf([scan_effect(i) for i in rasterize(letter)]).save(OUT / "scanned_letter.pdf")

    mixed = text_pdf(["\n\n".join(PARAGRAPHS_EN)])
    scanned_page = image_pdf([scan_effect(rasterize(text_pdf(["\n\n".join(PARAGRAPHS_ES)]))[0])])
    mixed.insert_pdf(scanned_page)
    mixed.save(OUT / "mixed_text_and_scan.pdf")

    # A long scanned document: slow enough to kill workers mid-processing in the chaos test.
    long_pages = [f"Page {n}\n\n" + "\n\n".join(random.sample(PARAGRAPHS_EN + PARAGRAPHS_ES, 5)) for n in range(1, 31)]
    image_pdf([scan_effect(i, angle=random.uniform(-1, 1)) for i in rasterize(text_pdf(long_pages), dpi=150)]).save(
        OUT / "scanned_report_30p.pdf"
    )

    # --- Images -------------------------------------------------------------------------
    invoice_img = rasterize(text_pdf([INVOICE], fontsize=12), dpi=220)[0].crop((100, 100, 1700, 1300))
    invoice_img.save(OUT / "invoice_scan.png")
    photo = scan_effect(invoice_img.convert("L"), angle=2.5, noise=30).convert("RGB")
    photo.save(OUT / "invoice_photo.jpg", quality=45)
    invoice_img.rotate(90, expand=True).save(OUT / "rotated_90.png")

    # --- Text formats -------------------------------------------------------------------
    (OUT / "notes_latin1.txt").write_bytes(("\n\n".join(PARAGRAPHS_ES)).encode("latin-1"))
    (OUT / "readme.md").write_text("# Sample\n\n" + "\n\n".join(f"- {p}" for p in PARAGRAPHS_EN), encoding="utf-8")
    (OUT / "prices.csv").write_text("item,qty,price\nOCR pages,1200,3600.00\nStorage GB,50,250.00\n", encoding="utf-8")
    (OUT / "article.html").write_text(
        "<!doctype html><html lang='en'><head><meta charset='utf-8'><title>Async pipelines</title>"
        "<meta name='description' content='Why queues matter'><style>p{color:red}</style>"
        "<script>console.log('ignored')</script></head><body><h1>Async pipelines</h1>"
        + "".join(f"<p>{p}</p>" for p in PARAGRAPHS_EN)
        + "<a href='https://example.com'>link</a></body></html>",
        encoding="utf-8",
    )
    (OUT / "memo.rtf").write_text(
        r"{\rtf1\ansi{\fonttbl\f0\fswiss Helvetica;}\f0\pard Memo: the pipeline is ready.\par "
        r"Contact ops@example.com before 2026-10-01.\par}",
        encoding="latin-1",
    )

    # --- Office -------------------------------------------------------------------------
    import docx
    import openpyxl
    from pptx import Presentation
    from pptx.util import Inches

    d = docx.Document()
    d.core_properties.title = "Quarterly report"
    d.core_properties.author = "Data Team"
    d.add_heading("Quarterly report", 0)
    for p in PARAGRAPHS_EN:
        d.add_paragraph(p)
    table = d.add_table(rows=3, cols=3)
    for r, row in enumerate([("Metric", "Q2", "Q3"), ("Documents", "12,400", "18,950"), ("Failures", "31", "12")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.add_paragraph("Conclusión: el sistema escaló sin pérdida de trabajos.")
    d.save(OUT / "report.docx")

    prs = Presentation()
    for title, body in [("Document Intelligence", PARAGRAPHS_EN[0]), ("Arquitectura", PARAGRAPHS_ES[2])]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = f"Speaker notes for {title}"
    prs.slides[0].shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(1)).text_frame.text = "Extra textbox"
    prs.save(OUT / "slides.pptx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Concept", "Amount (MXN)"])
    ws.append(["OCR", 3600])
    ws.append(["Storage", 250])
    ws.append(["Total", "=SUM(B2:B3)"])
    wb.create_sheet("Notes").append(["Prepared by billing@universidad.example.mx"])
    wb.save(OUT / "budget.xlsx")

    # --- Structured documents for LLM-ready Markdown ------------------------------------
    structured_pdf()
    scanned_structured()
    structured_docx(invoice_img)
    email_with_attachments()

    # --- Edge cases (should be rejected or fail gracefully) -----------------------------
    good = (OUT / "invoice_digital.pdf").read_bytes()
    (EDGE / "corrupted.pdf").write_bytes(good[: len(good) // 3])  # truncated
    (EDGE / "garbage.pdf").write_bytes(b"%PDF-1.7\n" + bytes(random.getrandbits(8) for _ in range(4000)))
    jpg = (OUT / "invoice_photo.jpg").read_bytes()
    (EDGE / "truncated.jpg").write_bytes(jpg[: len(jpg) // 4])
    (EDGE / "empty.txt").write_bytes(b"")
    (EDGE / "program.exe").write_bytes(b"MZ\x90\x00\x03\x00\x00\x00" + bytes(2000))
    with zipfile.ZipFile(EDGE / "archive.zip", "w") as zf:
        zf.writestr("a.txt", "hello")
    (EDGE / "text_named_as.png").write_text("This is plain text pretending to be a PNG image.", encoding="utf-8")
    (EDGE / "invoice_named_as.txt").write_bytes((OUT / "invoice_scan.png").read_bytes())
    enc = pymupdf.open(OUT / "invoice_digital.pdf")
    enc.save(EDGE / "encrypted.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
    blank = pymupdf.open()
    blank.new_page()
    blank.save(EDGE / "blank_page.pdf")
    Image.new("1", (20000, 20000), 1).save(EDGE / "huge_400mp.png", optimize=True)  # decompression bomb guard
    (EDGE / "binary_named_as.txt").write_bytes(bytes(random.getrandbits(8) for _ in range(5000)))

    for path in sorted(OUT.rglob("*")):
        if path.is_file():
            print(f"{path.relative_to(OUT)}  {path.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
